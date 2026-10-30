"""Extract and persist supported user-uploaded documents as indexed text sources."""
import csv
import io
import json
import re
import uuid
import zipfile
from itertools import chain
from xml.etree.ElementTree import ParseError
from pathlib import Path

from config import UPLOAD_DIR

SUPPORTED_EXTENSIONS = ("txt", "md", "csv", "json", "pdf", "docx", "xlsx")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 2_000_000


class UploadError(ValueError):
    """Raised when an uploaded document cannot be safely converted to text."""


def _checked_text(text):
    text = text.strip()
    if not text:
        raise UploadError("No readable text was found in the document.")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise UploadError("The extracted text exceeds the 2,000,000-character limit.")
    return text


def _join_text(parts):
    collected = []
    length = 0
    for part in parts:
        length += len(part)
        if length > MAX_EXTRACTED_CHARACTERS:
            raise UploadError("The extracted text exceeds the 2,000,000-character limit.")
        collected.append(part)
    return _checked_text("\n".join(collected))


def _decode(payload):
    try:
        return payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UploadError("Text files must use UTF-8 encoding.") from exc


def extract_text(filename, payload):
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS:
        raise UploadError(f"Unsupported file type: .{extension or '(none)'}")
    if not payload:
        raise UploadError("The uploaded file is empty.")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise UploadError("Each uploaded file must be 10 MB or smaller.")

    if extension in {"txt", "md"}:
        return _checked_text(_decode(payload))
    if extension == "csv":
        rows = csv.reader(io.StringIO(_decode(payload)))
        return _join_text(" | ".join(row) for row in rows)
    if extension == "json":
        try:
            data = json.loads(_decode(payload))
        except json.JSONDecodeError as exc:
            raise UploadError(f"Invalid JSON: {exc.msg}.") from exc
        return _checked_text(json.dumps(data, ensure_ascii=False, indent=2))
    if extension == "pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(payload))
            if reader.is_encrypted and not reader.decrypt(""):
                raise UploadError("Password-protected PDFs are not supported.")
            if len(reader.pages) > 500:
                raise UploadError("PDFs must contain 500 pages or fewer.")
            text = _join_text(page.extract_text() or "" for page in reader.pages)
        except UploadError:
            raise
        except (PdfReadError, ValueError, OSError) as exc:
            raise UploadError(f"Could not read the PDF ({type(exc).__name__}).") from exc
        return text
    if extension == "docx":
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
        from lxml.etree import XMLSyntaxError

        try:
            document = Document(io.BytesIO(payload))
            paragraph_parts = (paragraph.text for paragraph in document.paragraphs)
            table_rows = (
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            text = _join_text(chain(paragraph_parts, table_rows))
        except (ValueError, KeyError, OSError, PackageNotFoundError, XMLSyntaxError, zipfile.BadZipFile) as exc:
            raise UploadError(f"Could not read the Word document ({type(exc).__name__}).") from exc
        return text
    if extension == "xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)

            def sheet_parts():
                for sheet in workbook.worksheets:
                    yield f"Sheet: {sheet.title}"
                    for row in sheet.iter_rows(values_only=True):
                        yield " | ".join("" if value is None else str(value) for value in row)

            text = _join_text(sheet_parts())
        except (InvalidFileException, ValueError, KeyError, OSError, ParseError, zipfile.BadZipFile) as exc:
            raise UploadError(f"Could not read the spreadsheet ({type(exc).__name__}).") from exc
        finally:
            if workbook is not None:
                workbook.close()
        return text
    raise UploadError(f"Unsupported file type: .{extension}")


def save_uploaded_file(filename, payload, category="general", upload_dir=UPLOAD_DIR):
    category = category.casefold()
    if category not in {"general", "library", "transport"}:
        raise UploadError("Choose General, Library, or Transport as the document category.")
    text = extract_text(filename, payload)
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(filename).stem).strip("_-")[:60]
    stem = stem or "document"
    if stem.casefold().startswith(f"{category}_"):
        stem = stem[len(category) + 1:] or "document"
    upload_dir.mkdir(parents=True, exist_ok=True)
    destination = upload_dir / f"upload_{uuid.uuid4().hex}_{category}_{stem}.txt"
    destination.write_text(text, encoding="utf-8")
    return destination
