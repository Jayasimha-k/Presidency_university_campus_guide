from io import BytesIO
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from src.knowledge_ingestion import read_sources, source_name, sources
from src.uploaded_documents import extract_text, save_uploaded_file
from docx import Document
from openpyxl import Workbook


class KnowledgeIngestionTests(unittest.TestCase):
    def test_future_txt_route_files_are_discovered_automatically(self):
        with tempfile.TemporaryDirectory() as directory:
            knowledge_dir = Path(directory)
            upload_dir = knowledge_dir / "uploads"
            route_file = knowledge_dir / "future_routes.txt"
            route_file.write_text("ROUTE 999: Main Gate -> Gym", encoding="utf-8")
            (knowledge_dir / "notes.md").write_text("Not a source text file", encoding="utf-8")
            upload_dir.mkdir()
            upload_file = upload_dir / "library.txt"
            upload_file.write_text("Library rules", encoding="utf-8")

            with patch("src.knowledge_ingestion.KNOWLEDGE_DIR", knowledge_dir), \
                    patch("src.knowledge_ingestion.UPLOAD_DIR", upload_dir):
                found = sources()
                loaded = read_sources()
                names = [source_name(path) for path in found]

        self.assertEqual(names, ["future_routes.txt", "uploads/library.txt"])
        self.assertEqual(
            loaded,
            [(route_file, "ROUTE 999: Main Gate -> Gym"), (upload_file, "Library rules")],
        )

    def test_uploaded_csv_is_extracted_and_persisted_for_indexing(self):
        with tempfile.TemporaryDirectory() as directory:
            upload_dir = Path(directory) / "uploads"
            path = save_uploaded_file(
                "library_catalog.csv",
                b"Title,Location\nLaw books,M Block",
                category="Library",
                upload_dir=upload_dir,
            )
            text = path.read_text(encoding="utf-8")

            with patch("src.knowledge_ingestion.KNOWLEDGE_DIR", Path(directory) / "empty"), \
                    patch("src.knowledge_ingestion.UPLOAD_DIR", upload_dir):
                loaded = read_sources()
                name = source_name(path)

        self.assertEqual(path.suffix, ".txt")
        self.assertIn("_library_", path.name)
        self.assertEqual(text, "Title | Location\nLaw books | M Block")
        self.assertEqual(name, f"uploads/{path.name}")
        self.assertEqual(loaded[0][1], "Title | Location\nLaw books | M Block")

    def test_upload_rejects_unsupported_or_invalid_files(self):
        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            extract_text("library.exe", b"not a document")
        with self.assertRaisesRegex(ValueError, "Invalid JSON"):
            extract_text("library.json", b"{invalid}")

    def test_docx_and_xlsx_uploads_extract_searchable_text(self):
        word_stream = BytesIO()
        document = Document()
        document.add_paragraph("Law Library")
        document.save(word_stream)

        workbook_stream = BytesIO()
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Library"
        worksheet.append(["Facility", "Location"])
        worksheet.append(["Computers", "Technical Library"])
        workbook.save(workbook_stream)

        self.assertIn("Law Library", extract_text("library.docx", word_stream.getvalue()))
        spreadsheet_text = extract_text("library.xlsx", workbook_stream.getvalue())
        self.assertIn("Sheet: Library", spreadsheet_text)
        self.assertIn("Computers | Technical Library", spreadsheet_text)

    def test_pdf_upload_extracts_text(self):
        self.assertIn("Technical Library", extract_text("library.pdf", _make_text_pdf("Technical Library")))


def _make_text_pdf(text):
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(pdf)


if __name__ == "__main__":
    unittest.main()
